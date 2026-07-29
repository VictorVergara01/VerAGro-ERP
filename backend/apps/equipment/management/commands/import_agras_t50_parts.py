"""Importa los repuestos reales del DJI Agras T50 desde el Excel maestro.

Idempotente (upsert por SKU). Crea los 12 componentes (planos, bajo el modelo
T50), los productos (sku/part_number/name) y una ProductCompatibility por fila.
NO inventa números OEM: si la fila no trae 'Numero de Pieza', part_number queda vacío.
"""
import re

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from openpyxl import load_workbook

from apps.equipment.models import EquipmentModel, EquipmentComponent
from apps.inventory.models import Product, ProductCompatibility

SHEET = "T50 Todas las Partes"

# Excel (normalizado) -> (code, nombre visible)
CATEGORY_MAP = {
    "protector frontal": ("protector_frontal", "Protector frontal"),
    "m1-m2 brazos": ("brazos_m1_m2", "Brazos M1-M2"),
    "m3-m4 brazo": ("brazos_m3_m4", "Brazos M3-M4"),
    "chasis frontal": ("chasis_frontal", "Chasis frontal"),
    "chasis intermedio": ("chasis_intermedio", "Chasis intermedio"),
    "chasis trasero": ("chasis_trasero", "Chasis trasero"),
    "tren de aterrizaje": ("tren_aterrizaje", "Tren de aterrizaje"),
    "cableado de chasis": ("cableado_chasis", "Cableado de chasis"),
    "modulo de distribucion": ("modulo_distribucion", "Módulo de distribución"),
    "tanque de fumigacion": ("tanque_fumigacion", "Tanque de fumigación"),
    "sistema centrifugo": ("sistema_centrifugo", "Sistema centrífugo"),
    "helices": ("helices", "Hélices"),
}


def _slug(text):
    return re.sub(r"[^a-z0-9]+", "_", text.strip().lower()).strip("_")


class Command(BaseCommand):
    help = "Importa los repuestos reales del DJI Agras T50 desde el Excel maestro."

    def add_arguments(self, parser):
        parser.add_argument("xlsx", help="Ruta al archivo .xlsx")

    @transaction.atomic
    def handle(self, *args, **options):
        try:
            model = EquipmentModel.objects.get(model_code="T50")
        except EquipmentModel.DoesNotExist:
            raise CommandError(
                "No existe el modelo técnico T50. Corre primero seed_equipment_catalog."
            )

        wb = load_workbook(options["xlsx"], read_only=True, data_only=True)
        if SHEET not in wb.sheetnames:
            raise CommandError(f"El archivo no tiene la hoja '{SHEET}'.")
        ws = wb[SHEET]

        comp_cache = {}   # code -> EquipmentComponent
        n_comp = n_prod = n_compat = 0
        order = 0

        def get_component(categoria):
            nonlocal n_comp, order
            key = (categoria or "").strip().lower()
            if key in CATEGORY_MAP:
                code, name = CATEGORY_MAP[key]
            else:
                code, name = _slug(categoria or "sin_categoria"), (categoria or "Sin categoría")
                self.stdout.write(self.style.WARNING(f"Categoría no mapeada: {categoria!r} -> {code}"))
            if code not in comp_cache:
                comp, created = EquipmentComponent.objects.get_or_create(
                    equipment_model=model,
                    code=code,
                    defaults={
                        "name": name,
                        "component_type": EquipmentComponent.ComponentType.ASSEMBLY,
                        "diagram_key": code,
                        "sort_order": order,
                    },
                )
                order += 1
                if created:
                    n_comp += 1
                comp_cache[code] = comp
            return comp_cache[code]

        rows = ws.iter_rows(min_row=2, values_only=True)
        for row in rows:
            if not row or all(v in (None, "") for v in row[:5]):
                continue
            sku = (str(row[0]).strip() if row[0] is not None else "")
            part_number = (str(row[1]).strip() if row[1] is not None else "")
            name = (str(row[2]).strip() if row[2] is not None else "")
            categoria = row[4] if len(row) > 4 else None
            if not sku or not name:
                continue

            component = get_component(categoria)
            product, created = Product.objects.get_or_create(
                sku=sku,
                defaults={"name": name, "part_number": part_number},
            )
            if not created:
                changed = False
                if product.name != name:
                    product.name = name; changed = True
                if product.part_number != part_number:
                    product.part_number = part_number; changed = True
                if changed:
                    product.save(update_fields=["name", "part_number", "updated_at"])
            else:
                n_prod += 1

            _, c_created = ProductCompatibility.objects.get_or_create(
                product=product, equipment_model=model, component=component
            )
            if c_created:
                n_compat += 1

        self.stdout.write(self.style.SUCCESS(
            f"Import T50 OK: {n_comp} componentes, {n_prod} productos, {n_compat} compatibilidades nuevas."
        ))
